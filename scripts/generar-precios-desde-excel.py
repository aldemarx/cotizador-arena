#!/usr/bin/env python3
"""
Regenera src/data/precios-altos-de-sacta.json con el MODELO DEL EXCEL «COTIZADORES AGOSTO»
(hojas PRECIOS ALTOS y COTIZADOR ALTOS), que no cobra interés:

  precio con dscto  = tarifa_m2 × metraje + 1000      (H)   = contado + inicial
  precio al contado = precio con dscto − 2000         (I)
  precio total      = precio con dscto + 7000         (G, «precio real», va tachado)
  inicial           = 2000 en dos partes (1000 + 1000)
  cuota fija        = (precio con dscto − inicial) / 80 = contado / 80

Las cuotas llevan DECIMALES, igual que el Excel (se redondean a céntimos). Ascendente y descendente no
existen en el Excel: se derivan con la misma variación anual de ±5 % y el mismo total que la fija
(con tasa 0), para no perder la pantalla «Opciones de cuota». Pendiente de confirmar con ventas.

Metraje, ubicación, etapa y tarifa se toman del JSON actual (coinciden con el Excel en los 1,112
lotes); los precios se recalculan y se comparan con los valores calculados del Excel.

Uso:  python scripts/generar-precios-desde-excel.py "icons/COTIZADORES  AGOSTO.xlsx"
Vuelta atrás: git checkout checkpoint-precios-modelo-interes -- src/data/precios-altos-de-sacta.json
              (o copiar src/data/precios-altos-de-sacta.modelo-interes.json sobre el JSON)
"""
import json, math, sys
import openpyxl

RUTA_XLSX = sys.argv[1]
JSON_PATH = "src/data/precios-altos-de-sacta.json"
RESPALDO = "src/data/precios-altos-de-sacta.modelo-interes.json"

base = json.load(open(RESPALDO, encoding="utf-8"))  # mismo metraje/ubicación/tarifa que el Excel
P = base["parametros"]
N, MESES, TRAMOS = P["numero_cuotas"], P["meses_por_tramo"], P["tramos"]
INICIAL, M1 = 2000, 1000
M2 = INICIAL - M1
tamanos = [MESES if k < TRAMOS - 1 else N - MESES * (TRAMOS - 1) for k in range(TRAMOS)]
tramo_de = lambda m: (m - 1) // MESES

# Medio céntimo sube (como Excel); la tolerancia evita que el ruido de coma flotante lo baje.
centimos = lambda x: math.floor(x * 100 + 0.5 + 1e-7) / 100
def num(x):
    x = round(x, 2)
    return int(x) if x == int(x) else x

def plan(pv, g):
    den = sum((1 + g) ** tramo_de(m) for m in range(1, N + 1))
    C = pv / den
    cuotas = [num(centimos(C * (1 + g) ** k)) for k in range(TRAMOS)]
    total = num(sum(c * t for c, t in zip(cuotas, tamanos)))
    return {"cuotas_por_tramo": cuotas, "total_cuotas": total, "precio_financiado": num(INICIAL + total)}

# Valores calculados por el Excel, para contrastar
ws = openpyxl.load_workbook(RUTA_XLSX, data_only=True)["PRECIOS ALTOS"]
excel = {}
for r in range(7, ws.max_row + 1):
    k = ws.cell(r, 1).value
    if k:
        excel[k] = {"G": ws.cell(r, 7).value, "H": ws.cell(r, 8).value, "I": ws.cell(r, 9).value}

lotes, sin_comparar, dif = [], [], []
for l in base["lotes"]:
    H = round(l["tarifa_m2"] * l["metraje_m2"] + 1000, 2)
    contado, lista = round(H - INICIAL, 2), round(H + 7000, 2)
    e = excel[l["lote"]]
    if all(isinstance(e[c], (int, float)) for c in "GHI"):
        if max(abs(e["G"] - lista), abs(e["H"] - H), abs(e["I"] - contado)) > 0.01:
            dif.append(l["lote"])
    else:
        sin_comparar.append(l["lote"])  # p. ej. 5-574: el Excel trae el metraje como texto y da #VALUE!
    lotes.append({
        "lote": l["lote"], "etapa": l["etapa"], "mz": l["mz"], "lt": l["lt"],
        "metraje_m2": l["metraje_m2"], "ubicacion": l["ubicacion"], "tarifa_m2": l["tarifa_m2"],
        "precio_lista": num(lista),
        "precio_con_dscto": num(H),
        "precio_contado": num(contado),
        "inicial": INICIAL, "inicial_mes_1": M1, "inicial_mes_2": M2,
        "planes": {
            "fija": plan(contado, 0),
            "ascendente": plan(contado, P["variacion_anual_ascendente"]),
            "descendente": plan(contado, P["variacion_anual_descendente"]),
        },
    })

salida = {**base, "version": "2026-10-03-excel-agosto", "lotes": lotes}
salida["parametros"] = {
    **P,
    "tasa_mensual": 0,
    "modelo": "excel-agosto-2026",
    "en_prueba": "Precios del Excel «COTIZADORES AGOSTO» (sin interés, inicial fija S/ 2,000). Vuelta atrás: git checkout checkpoint-precios-modelo-interes -- src/data/precios-altos-de-sacta.json",
}
salida["parametros"].pop("descuento_lista", None)
salida["parametros"]["recargo_lista"] = 7000
salida["parametros"]["calendario"] = "inicial 1000 + 1000; 80 cuotas; cuota fija = (precio con dscto − inicial) / 80, con decimales como el Excel"
# Misma indentación (1 espacio) y sin salto final que el archivo original.
open(JSON_PATH, "w", encoding="utf-8", newline="\n").write(json.dumps(salida, indent=1, ensure_ascii=False))
print(f"Escrito {JSON_PATH}: {len(lotes)} lotes · versión {salida['version']}")
print(f"Contraste con el Excel: {len(excel)-len(sin_comparar)-len(dif)} coinciden, {len(dif)} difieren {dif[:5]}, sin comparar {sin_comparar}")
