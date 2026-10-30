"""Aplica al JSON de lotes los dos Excels de promoción de octubre 2026 (mandan sobre lo cargado antes):
  · PRECIOS PROMOCION ARENAS 2905.xlsx  → R1–R15: metraje, precios y promoción Standard / Preferencial.
  · PRECIOS MZ 22-23-24.xlsx            → R22–R24: «precio venta» = precio final (el de lanzamiento se omite); sin promoción.
Los lotes con promoción se financian solo a 60 cuotas (lo fuerza la interfaz con `promocion != null`).

Uso: python -I scripts/aplicar-excels-promocion.py <promocion.xlsx> <mz-22-23-24.xlsx> [--aplicar]
Sin --aplicar solo muestra los cambios. Después: node scripts/migrar-lotes-a-supabase.mjs (sube el JSON a Supabase).
"""
import json, sys
import openpyxl

promo_xlsx, mz_xlsx = sys.argv[1], sys.argv[2]
aplicar = "--aplicar" in sys.argv
RUTA = "src/data/arenas-lotes.json"
# R9-12: el Excel trae «promoción contado» 4,300 (es la constante del descuento, no un precio): se conserva la promoción ya cargada.
CONSERVAR_PROMO = {"R9-12"}

datos = json.load(open(RUTA, encoding="utf8"))
J = {l["lote"]: l for l in datos["lotes"]}
cambios = []


def r2(x):
    return round(x + 1e-9, 2)


def poner(l, campo, valor):
    if l.get(campo) != valor:
        cambios.append((l["lote"], campo, l.get(campo), valor))
        l[campo] = valor


def cuotas(l):
    saldo = l["precio_lista"] - l["inicial"]
    poner(l, "cuota_60", r2(saldo / 60))
    poner(l, "cuota_80", r2(saldo / 80))


# ---- PRECIOS PROMOCION ARENAS 2905 (R1–R15) ----
ws = openpyxl.load_workbook(promo_xlsx, data_only=True).worksheets[0]
for r in ws.iter_rows(values_only=True):
    if not (isinstance(r[0], str) and r[0].startswith("R") and "-" in r[0]):
        continue
    l = J[r[0].strip()]
    real, dscto, p_cont, p_fin, contado, inicial, tarifa = r[7], r[8], r[9], r[10], r[11], r[12], r[14]
    if abs(dscto - (contado + inicial)) > 0.01:
        # R15-12 trae inicial 2,247 pero su contado es dscto - 2,000: se toma la inicial de 2,000 del resto de lotes.
        print("AVISO", r[0], "inicial", inicial, "no cuadra con dscto - contado; se usa", dscto - contado)
        inicial = dscto - contado
    poner(l, "metraje_m2", float(r[4]))
    poner(l, "tarifa_m2", float(tarifa))
    poner(l, "precio_lista", r2(dscto))
    poner(l, "precio_contado", r2(contado))
    poner(l, "inicial", float(inicial))
    cuotas(l)
    if l["lote"] in CONSERVAR_PROMO:
        poner(l, "pendientes", ["promoción: el Excel trae «promoción contado» 4,300 (error de digitación); se conserva la ya cargada"])
        continue
    poner(l, "pendientes", [])
    if p_cont is None:
        continue  # sin promoción en el Excel: queda como estaba (R5–R7, R11–R15 no la tienen)
    poner(l, "promocion", {
        "tipo": "standard" if p_fin is not None else "preferencial",
        "precio_real": r2(real),
        "precio_con_dscto": r2(dscto),
        "promo_contado": r2(p_cont),
        "promo_financiado": r2(p_fin) if p_fin is not None else None,
    })

# ---- PRECIOS MZ 22-23-24 (R22-R24) ----
# El «precio venta» es el precio final; el «precio lanzamiento» se omite. Son lotes normales: sin promoción, plazo libre.
ws = openpyxl.load_workbook(mz_xlsx, data_only=True).worksheets[0]
for r in ws.iter_rows(values_only=True):
    i = next((k for k, c in enumerate(r) if isinstance(c, str) and c.startswith("R") and "-" in c), None)
    if i is None:
        continue
    l = J[r[i].strip()]
    # «precio venta»: la cifra que va seguida de la inicial (2,000) y del contado (venta - inicial); justo antes va su valor por m2
    k = next(k for k in range(i + 5, len(r) - 2) if r[k + 1] == 2000 and abs(r[k + 2] - (r[k] - 2000)) < 0.01)
    venta, val = r[k], r[k - 1]
    poner(l, "metraje_m2", float(r[i + 1]))
    poner(l, "tarifa_m2", float(val))
    poner(l, "precio_lista", r2(venta))
    poner(l, "precio_contado", r2(venta - 2000))
    poner(l, "inicial", 2000.0)
    cuotas(l)
    poner(l, "pendientes", [])
    poner(l, "promocion", None)

con_promo = sum(1 for l in datos["lotes"] if l["promocion"])
por_campo = {}
for _, campo, _, _ in cambios:
    por_campo[campo] = por_campo.get(campo, 0) + 1
print("cambios por campo:", por_campo)
print("lotes con promoción tras aplicar:", con_promo, "de", len(datos["lotes"]))
for c in cambios:
    if c[1] == "precio_lista":
        print("  ", c)
if aplicar:
    with open(RUTA, "w", encoding="utf8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print("JSON actualizado.")
