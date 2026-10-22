"""Consolida los dos Excel de Arenas en un solo dataset.

Base: 'ARENAS Y AURA' (completo, solo la hoja ARENAS: 654 lotes). Sobre los lotes repetidos se aplican las condiciones
de 'PRECIOS PROMOCION ARENAS' (promoción especial), salvo los que tienen dudas abiertas
(ver docs/pendientes-arenas.md): esos se quedan con los valores del completo.

Uso: python -I scripts/consolidar-arenas.py <completo.xlsx> <promocion.xlsx>
Genera: src/data/arenas-lotes.json y datos/ARENAS-consolidado.xlsx
"""
import json, sys, collections
import openpyxl

completo_xlsx, promo_xlsx = sys.argv[1], sys.argv[2]

# Lotes con dudas abiertas: se queda lo del completo (ver pendientes).
PEND_METRAJE = {"R1-1", "R1-28", "R2-1", "R2-12"}   # completo trae 148 m² (= tarifa); promo trae otro
PEND_UBICACION = {"R4-16", "R4-17"}                 # completo: ESQUINA - AV PRINCIPAL; promo: ESQUINA
PEND_INCOMPLETO_PROMO = {"R8", "R10"}               # promo Standard sin precios completos en esas manzanas


def num(x):
    return round(float(x), 2) if isinstance(x, (int, float)) else None


full = openpyxl.load_workbook(completo_xlsx, data_only=True)
promo = openpyxl.load_workbook(promo_xlsx, data_only=True)

lotes = {}
for ws in full:
    if ws.title != "ARENAS":  # la hoja AURA es otro proyecto: fuera de alcance
        continue
    proyecto = "arenas"
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[2]:
            continue
        cod = r[2]
        mz = cod.split("-")[0]
        contado = num(r[8])
        lotes[cod] = {
            "lote": cod, "proyecto": proyecto, "mz": mz, "lt": r[1],
            "metraje_m2": num(r[3]), "ubicacion": str(r[5]).replace("�", "Á").strip(),
            "tarifa_m2": num(r[6]), "precio_lista": num(r[7]), "precio_contado": contado,
            "inicial": num(r[9]),
            "cuota_60": num(r[10]), "cuota_80": num(r[12]),
            "promocion": None, "pendientes": [],
        }

aplicadas = 0
for ws in promo:
    tipo = "preferencial" if "Preferencial" in ws.title else "standard"
    for r in ws.iter_rows(min_row=7, values_only=True):
        cod = r[0]
        if not cod:
            continue
        l = lotes[cod]  # todos existen en el completo (auditado)
        if cod in PEND_METRAJE:
            l["pendientes"].append(f"metraje: completo {l['metraje_m2']} m² vs promoción {num(r[4])} m²")
        if cod in PEND_UBICACION:
            l["pendientes"].append(f"ubicación: completo '{l['ubicacion']}' vs promoción '{str(r[5]).strip()}'")
        if l["mz"] in PEND_INCOMPLETO_PROMO and tipo == "standard":
            l["pendientes"].append("promoción Standard incompleta en el Excel de promoción")
        if l["pendientes"]:
            continue
        # Verificación: el contado de la promoción debe cuadrar con el del completo
        assert num(r[10]) == l["precio_contado"], (cod, r[10], l["precio_contado"])
        l["promocion"] = {
            "tipo": tipo,
            "precio_real": num(r[6]),
            "precio_con_dscto": num(r[7]),
            "promo_contado": num(r[8]),
            "promo_financiado": num(r[9]),  # None = 'No aplica'
        }
        aplicadas += 1

out = sorted(lotes.values(), key=lambda l: (l["proyecto"], l["mz"], l["lt"]))
res = collections.Counter(l["proyecto"] for l in out)
print("lotes:", dict(res), "| con promoción aplicada:", aplicadas,
      "| con pendientes:", sum(1 for l in out if l["pendientes"]))

with open("src/data/arenas-lotes.json", "w", encoding="utf-8") as f:
    json.dump({"proyecto": "Arenas del Valle", "lotes": out}, f, ensure_ascii=False, indent=1)

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "LOTES"
cab = ["Lote", "Proyecto", "Mz", "Lt", "Metraje m²", "Ubicación", "Tarifa m²", "Precio lista", "Precio contado",
       "Inicial", "Cuota 60", "Cuota 80", "Promo", "Precio real (promo)", "Precio con dscto (promo)",
       "Promo contado", "Promo financiado", "Pendientes"]
ws.append(cab)
for l in out:
    p = l["promocion"] or {}
    ws.append([l["lote"], l["proyecto"], l["mz"], l["lt"], l["metraje_m2"], l["ubicacion"], l["tarifa_m2"],
               l["precio_lista"], l["precio_contado"], l["inicial"], l["cuota_60"], l["cuota_80"],
               p.get("tipo"), p.get("precio_real"), p.get("precio_con_dscto"), p.get("promo_contado"),
               p.get("promo_financiado") if p else None, "; ".join(l["pendientes"])])
ws.freeze_panes = "B2"
for col, w in zip("ABCDEFGHIJKLMNOPQR", [9, 9, 6, 5, 11, 28, 10, 12, 13, 8, 10, 10, 13, 16, 18, 14, 16, 60]):
    ws.column_dimensions[col].width = w
wb.save("datos/ARENAS-consolidado.xlsx")
